#!/usr/bin/env python3
"""Generate an Anki .apkg deck from a Feishu/Lark spreadsheet OR a local .xlsx.

Two data sources:
  --url  <FEISHU_SHEET_URL>
      Reads via lark-cli (sheets +cells-get) and downloads embedded in-cell
      images via docs +media-download (no workbook-export scope needed).
  --xlsx <PATH>
      Reads values with openpyxl and extracts embedded images straight from
      the xlsx package (xl/drawings + xl/media), fully offline.

CSV cannot carry images, which is why this tool works with Feishu sheets and
xlsx files rather than CSV exports.

Header convention (row 1, matched case-insensitively, Chinese/English):
  Deck* / 牌组*      one or more columns, joined with "::" into nested decks
  Front / 正面       card front (text and/or embedded image)
  Back / 背面        card back
  Note Type / 卡片类型  "Basic"(default) or "Cloze"/"填空"
  Tags* / 标签*      tags, split on whitespace / comma / semicolon
  Fields*_x / 字段*_x  extra Anki fields; display name is the part after "_"

Usage:
  python3 feishu_to_anki.py --url  <SHEET_URL> -o out.apkg [--sheet-name N]
  python3 feishu_to_anki.py --xlsx <FILE.xlsx> -o out.apkg [--sheet-name N]
"""

import argparse
import glob
import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET

try:
    import genanki
except ImportError:
    sys.stderr.write(
        "ERROR: genanki is not installed. Run:\n"
        "  python3 -m pip install --user genanki openpyxl\n"
    )
    sys.exit(2)

MODEL_BASIC_ID = 1726354981
MODEL_CLOZE_ID = 1726354982

CSS = """
.card { font-family: -apple-system, "PingFang SC", "Microsoft YaHei", sans-serif;
        font-size: 16px; text-align: center; color: #222; background: #fff; }
img { max-width: 90%; max-height: 420px; border-radius: 8px; }
hr { border: none; border-top: 1px solid #ccc; margin: 14px 0; }
.extra { margin-top: 6px; font-size: 14px; color: #555; text-align: left; }
.extra .lbl { font-weight: bold; color: #333; }
#answer { text-align: left; line-height: 1.6; white-space: normal; }
"""

LOCAL_PREFIX = "file://"
NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


# ================================================================ lark-cli

def resolve_lark_cli():
    env = os.environ.get("LARK_CLI")
    if env and os.path.exists(env):
        return env
    found = shutil.which("lark-cli")
    if found:
        return found
    matches = sorted(glob.glob(os.path.expanduser(
        "~/.trae-cn/plugins/trae-remote-official/lark/*/bin/lark-cli")))
    if matches:
        return matches[-1]
    sys.stderr.write("ERROR: lark-cli not found. Set LARK_CLI or install the Lark plugin.\n")
    sys.exit(2)


def lark_json(cli, args, cwd=None):
    proc = subprocess.run([cli] + args, capture_output=True, text=True, cwd=cwd)
    if proc.returncode != 0:
        detail = proc.stderr.strip() or proc.stdout.strip()
        raise RuntimeError(f"lark-cli failed ({' '.join(args[:2])}):\n{detail}")
    try:
        payload = json.loads(proc.stdout)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"non-JSON output from lark-cli: {exc}\n{proc.stdout[:500]}")
    if not payload.get("ok"):
        raise RuntimeError(
            f"lark-cli error: {json.dumps(payload.get('error', payload), ensure_ascii=False)}")
    return payload["data"]


def col_to_index(letters):
    idx = 0
    for ch in letters:
        idx = idx * 26 + (ord(ch.upper()) - ord("A") + 1)
    return idx - 1


def index_to_col(idx):
    s = ""
    idx += 1
    while idx:
        idx, r = divmod(idx - 1, 26)
        s = chr(65 + r) + s
    return s


# ================================================================ Feishu

def fetch_feishu(cli, url, sheet_name):
    wb = lark_json(cli, ["sheets", "+workbook-info", "--url", url])
    title = wb.get("title") or wb.get("spreadsheet", {}).get("title") or "Anki"
    sheets = wb.get("sheets") or []
    if not sheets:
        raise RuntimeError("workbook has no sheets")
    chosen = None
    if sheet_name:
        for s in sheets:
            if s.get("sheet_name") == sheet_name or s.get("sheet_id") == sheet_name:
                chosen = s
                break
        if chosen is None:
            raise RuntimeError(
                f"sheet '{sheet_name}' not found; available: "
                + ", ".join(s.get("sheet_name", "?") for s in sheets))
    else:
        chosen = sheets[0]
    sid = chosen["sheet_id"]
    name = chosen.get("sheet_name") or sid

    csv_data = lark_json(cli, ["sheets", "+csv-get", "--url", url, "--sheet-id", sid])
    region = csv_data.get("current_region") or "A1:A1"
    cells_data = lark_json(cli, ["sheets", "+cells-get", "--url", url,
                                 "--sheet-id", sid, "--range", region,
                                 "--include", "value"])
    grid = {}
    for rng in cells_data.get("ranges", []):
        col_ids = [col_to_index(c) for c in rng["col_indices"]]
        for ri, row_cells in zip(rng["row_indices"], rng["cells"]):
            row_map = grid.setdefault(ri, {})
            for ci, cell in zip(col_ids, row_cells):
                if cell:
                    row_map[ci] = cell
    return title, sid, name, grid, None  # last item = temp dir to clean later


# ================================================================ xlsx

def _rels_map(rels_xml_bytes):
    """Return {rId: target_path} for a *.rels package part."""
    root = ET.fromstring(rels_xml_bytes)
    out = {}
    for rel in root.findall("rel:Relationship", NS):
        out[rel.get("Id")] = rel.get("Target")
    return out


def _norm_target(base_part_dir, target):
    """Resolve a relationship target (e.g. ../drawings/drawing1.xml) to a zip path."""
    if target.startswith("/"):
        return target.lstrip("/")
    combined = os.path.normpath(os.path.join(base_part_dir, target)).replace("\\", "/")
    return combined


def _rels_path(part):
    """Zip path of the .rels sidecar for a package part."""
    return f"{os.path.dirname(part)}/_rels/{os.path.basename(part)}.rels"


def _xlsx_image_anchors(zf, sheet_part):
    """Map {(row_1based, col_0based): [media_zip_path, ...]} for one worksheet."""
    anchors = {}
    sheet_dir = os.path.dirname(sheet_part)  # xl/worksheets
    try:
        sheet_rels = _rels_map(zf.read(_rels_path(sheet_part)))
    except KeyError:
        return anchors
    drawing_targets = [
        _norm_target(sheet_dir, t) for rid, t in sheet_rels.items()
        if t.endswith(".xml") and "drawings/drawing" in t
    ]
    for drawing_part in drawing_targets:
        draw_dir = os.path.dirname(drawing_part)  # xl/drawings
        try:
            draw_rels = _rels_map(zf.read(_rels_path(drawing_part)))
        except KeyError:
            draw_rels = {}
        root = ET.fromstring(zf.read(drawing_part))
        for anchor in list(root):
            frm = anchor.find("xdr:from", NS)
            if frm is None:
                continue
            row_el = frm.find("xdr:row", NS)
            col_el = frm.find("xdr:col", NS)
            if row_el is None or col_el is None:
                continue
            row0, col0 = int(row_el.text), int(col_el.text)
            for blip in anchor.iter(f"{{{NS['a']}}}blip"):
                rid = blip.get(f"{{{NS['r']}}}embed")
                target = draw_rels.get(rid)
                if not target:
                    continue
                media_path = _norm_target(draw_dir, target)
                anchors.setdefault((row0 + 1, col0), []).append(media_path)
    return anchors


def fetch_xlsx(path, sheet_name, media_tmp):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("openpyxl is not installed. Run: python3 -m pip install --user openpyxl")

    wb = load_workbook(path, data_only=True, read_only=True)
    names = wb.sheetnames
    if not names:
        raise RuntimeError("xlsx has no worksheets")
    if sheet_name:
        if sheet_name not in names:
            raise RuntimeError(f"sheet '{sheet_name}' not found; available: {', '.join(names)}")
        ws = wb[sheet_name]
        ws_index = names.index(sheet_name)
    else:
        ws = wb[names[0]]
        ws_index = 0

    grid = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            grid.setdefault(cell.row, {})[cell.column - 1] = {"value": str(cell.value)}

    title = os.path.splitext(os.path.basename(path))[0]
    images = {}
    with zipfile.ZipFile(path) as zf:
        wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
        sheets_xml = wb_root.find("main:sheets", NS)
        rids = [s.get(f"{{{NS['r']}}}id") for s in list(sheets_xml)]
        wb_rels = _rels_map(zf.read("xl/_rels/workbook.xml.rels"))
        sheet_part = _norm_target("xl", wb_rels[rids[ws_index]])
        raw_anchors = _xlsx_image_anchors(zf, sheet_part)

        for (row1, col0), media_paths in raw_anchors.items():
            rich = []
            for media_path in media_paths:
                ext = os.path.splitext(media_path)[1] or ".img"
                out_name = f"{hashlib.md5(media_path.encode()).hexdigest()[:12]}{ext}"
                out_path = os.path.join(media_tmp, out_name)
                if not os.path.exists(out_path):
                    with open(out_path, "wb") as fh:
                        fh.write(zf.read(media_path))
                rich.append({"type": "embed-image",
                             "image_token": LOCAL_PREFIX + out_path,
                             "image_name": os.path.basename(media_path)})
            existing = grid.setdefault(row1, {}).get(col0, {})
            if "value" in existing:
                existing["rich_text"] = rich
            else:
                existing["rich_text"] = rich
                grid[row1][col0] = existing
            images[(row1, col0)] = rich
    wb.close()
    return title, os.path.abspath(path), ws.title, grid, media_tmp


# ================================================================ cells

def cell_text(cell):
    if not cell:
        return ""
    if cell.get("value") is not None:
        return str(cell["value"])
    parts = []
    for seg in cell.get("rich_text", []) or []:
        if seg.get("type") != "embed-image" and seg.get("text"):
            parts.append(seg["text"])
    return "".join(parts)


def cell_images(cell):
    if not cell:
        return []
    return [seg["image_token"] for seg in cell.get("rich_text", []) or []
            if seg.get("type") == "embed-image" and seg.get("image_token")]


# ================================================================ headers

RE_DECK = re.compile(r"^(deck|牌组)", re.I)
RE_FRONT = re.compile(r"^(front|正面)", re.I)
RE_BACK = re.compile(r"^(back|背面)", re.I)
RE_TYPE = re.compile(r"(note\s*type|卡片类型)", re.I)
RE_TAGS = re.compile(r"^(tags?|标签)", re.I)
RE_FIELD = re.compile(r"^(?:fields?\d*|字段\d*)\s*[_:\s：-]*(.*)$", re.I)


def classify_headers(headers):
    plan = {"decks": [], "front": None, "back": None,
            "type": None, "tags": [], "extras": []}
    for ci in sorted(headers):
        h = (headers[ci] or "").strip()
        if not h:
            continue
        if RE_DECK.match(h):
            plan["decks"].append(ci)
        elif RE_FRONT.match(h):
            plan["front"] = ci
        elif RE_BACK.match(h):
            plan["back"] = ci
        elif RE_TYPE.search(h):
            plan["type"] = ci
        elif RE_TAGS.match(h):
            plan["tags"].append(ci)
        else:
            m = RE_FIELD.match(h)
            if m:
                display = (m.group(1) or h).strip() or h
                plan["extras"].append((ci, display))
    return plan


def text_to_html(text):
    text = (text or "").strip()
    return html.escape(text).replace("\n", "<br>") if text else ""


def split_tags(raw):
    out = []
    for t in re.split(r"[\s,，;；]+", raw or ""):
        t = t.strip().lstrip("#").replace(" ", "_")
        if t and t not in out:
            out.append(t)
    return out


# ================================================================ main

def main():
    ap = argparse.ArgumentParser(description="Build .apkg from a Feishu sheet or xlsx file")
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--url", help="Feishu spreadsheet URL")
    src.add_argument("--xlsx", help="local .xlsx path (embedded images supported)")
    ap.add_argument("--sheet-name", help="sub-sheet title; default: first sheet")
    ap.add_argument("-o", "--output", help="output .apkg path")
    ap.add_argument("--deck", help="fallback deck name when Deck columns are empty")
    ap.add_argument("--media-dir", help="directory for prepared images (default: temp)")
    ap.add_argument("--keep-media", action="store_true", help="keep prepared images")
    args = ap.parse_args()

    cleanup_dirs = []
    cli = None
    if args.url:
        cli = resolve_lark_cli()
        wb_title, source_key, sheet_name, grid, tmp_dir = \
            fetch_feishu(cli, args.url, args.sheet_name)
    else:
        if not os.path.exists(args.xlsx):
            raise RuntimeError(f"xlsx not found: {args.xlsx}")
        xlsx_tmp = tempfile.mkdtemp(prefix="xlsx_media_")
        cleanup_dirs.append(xlsx_tmp)
        wb_title, source_key, sheet_name, grid, _ = \
            fetch_xlsx(args.xlsx, args.sheet_name, xlsx_tmp)

    if not grid:
        raise RuntimeError("sheet is empty")

    header_row = min(grid)
    headers = {ci: cell_text(c) for ci, c in grid[header_row].items()}
    plan = classify_headers(headers)
    if plan["front"] is None and plan["back"] is None:
        raise RuntimeError("no Front/正面 or Back/背面 column found in header row")

    extra_names = []
    for _, name in plan["extras"]:
        if name not in extra_names:
            extra_names.append(name)

    media_dir = args.media_dir or tempfile.mkdtemp(prefix="anki_media_")
    os.makedirs(media_dir, exist_ok=True)
    if not args.media_dir:
        cleanup_dirs.append(media_dir)
    media_files = []

    def acquire_image(token, row, ci):
        stem = f"r{row}_{index_to_col(ci)}_{hashlib.md5(token.encode()).hexdigest()[:8]}"
        if token.startswith(LOCAL_PREFIX):
            src_path = token[len(LOCAL_PREFIX):]
            ext = os.path.splitext(src_path)[1]
            dst = os.path.join(media_dir, stem + ext)
            shutil.copyfile(src_path, dst)
        else:
            data = lark_json(cli,
                             ["docs", "+media-download", "--token", token,
                              "--output", stem], cwd=media_dir)
            dst = data.get("saved_path") or os.path.join(media_dir, stem)
            if not os.path.exists(dst):
                hits = glob.glob(os.path.join(media_dir, stem + ".*"))
                dst = hits[0] if hits else dst
        media_files.append(dst)
        return os.path.basename(dst)

    def cell_html(row, ci):
        cell = grid.get(row, {}).get(ci)
        bits = []
        txt = text_to_html(cell_text(cell))
        if txt:
            bits.append(txt)
        for token in cell_images(cell):
            bits.append(f'<img src="{acquire_image(token, row, ci)}">')
        return "<br>".join(bits)

    # ---- models
    basic_fields = [{"name": "Front"}, {"name": "Back"}] + \
                   [{"name": n} for n in extra_names]
    extra_blocks = "\n".join(
        f'{{{{#{n}}}}}<div class="extra"><span class="lbl">{html.escape(n)}：</span>'
        f'{{{{{n}}}}}</div>{{{{/{n}}}}}'
        for n in extra_names
    )
    basic_model = genanki.Model(
        MODEL_BASIC_ID, "Sheet2Anki Basic",
        fields=basic_fields,
        templates=[{
            "name": "Card 1",
            "qfmt": '<div class="front">{{Front}}</div>',
            "afmt": '{{FrontSide}}<hr id="answer">{{Back}}' + "\n" + extra_blocks,
        }],
        css=CSS)

    cloze_fields = [{"name": "Text"}, {"name": "Back"}] + \
                   [{"name": n} for n in extra_names]
    cloze_model = genanki.Model(
        MODEL_CLOZE_ID, "Sheet2Anki Cloze",
        fields=cloze_fields,
        templates=[{
            "name": "Cloze",
            "qfmt": '{{cloze:Text}}',
            "afmt": '{{cloze:Text}}<hr id="answer">{{Back}}' + "\n" + extra_blocks,
        }],
        css=CSS)

    decks = {}

    def get_deck(row):
        parts = [cell_text(grid.get(row, {}).get(ci)).strip()
                 for ci in plan["decks"]]
        parts = [p for p in parts if p]
        path = "::".join(parts) if parts else (args.deck or wb_title or "Sheet2Anki")
        if path not in decks:
            deck_id = int(hashlib.md5(path.encode("utf-8")).hexdigest()[:8], 16)
            decks[path] = genanki.Deck(deck_id=deck_id, name=path)
        return decks[path]

    notes = skipped = cloze_rows = 0
    for row in sorted(r for r in grid if r > header_row):
        front = cell_html(row, plan["front"]) if plan["front"] is not None else ""
        back = cell_html(row, plan["back"]) if plan["back"] is not None else ""
        if not front and not back:
            skipped += 1
            continue
        extras_vals = [cell_html(row, ci) for ci, _ in plan["extras"]]
        tags = []
        for ci in plan["tags"]:
            tags.extend(split_tags(cell_text(grid.get(row, {}).get(ci))))
        note_type = cell_text(grid.get(row, {}).get(plan["type"])).lower() \
            if plan["type"] is not None else ""
        is_cloze = ("cloze" in note_type) or ("填空" in note_type)
        if is_cloze:
            cloze_rows += 1
            note = genanki.Note(
                model=cloze_model,
                fields=[front or back, back if front else ""] + extras_vals,
                tags=tags, guid=genanki.guid_for(source_key, row))
        else:
            note = genanki.Note(
                model=basic_model,
                fields=[front, back] + extras_vals,
                tags=tags, guid=genanki.guid_for(source_key, row))
        get_deck(row).add_note(note)
        notes += 1

    if notes == 0:
        raise RuntimeError("no card rows found")

    output = args.output or os.path.join(os.getcwd(), f"{wb_title}.apkg")
    package = genanki.Package(list(decks.values()))
    package.media_files = media_files
    package.write_to_file(output)

    if not args.keep_media:
        for d in cleanup_dirs:
            shutil.rmtree(d, ignore_errors=True)

    print(json.dumps({
        "ok": True,
        "output": os.path.abspath(output),
        "source": "feishu" if args.url else "xlsx",
        "spreadsheet": wb_title,
        "sheet": sheet_name,
        "notes": notes,
        "cloze_notes": cloze_rows,
        "skipped_empty_rows": skipped,
        "decks": sorted(decks.keys()),
        "images": len(media_files),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False),
              file=sys.stderr)
        sys.exit(1)
